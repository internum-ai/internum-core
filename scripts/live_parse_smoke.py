import io
import json
import os
import shutil
import subprocess
import tempfile
from collections.abc import Iterable
from pathlib import Path
from typing import Any

from api.main import create_app
from fastapi.testclient import TestClient


def main() -> int:
    if os.getenv("RUN_LIVE_OPENROUTER_SMOKE") != "true":
        print("Live OpenRouter smoke skipped: set RUN_LIVE_OPENROUTER_SMOKE=true to enable.")
        return 0

    api_key = os.getenv("CORE_LIVE_SMOKE_API_KEY")
    if not api_key:
        print("Live OpenRouter smoke skipped: CORE_LIVE_SMOKE_API_KEY is required.")
        return 1

    schema = {
        "type": "object",
        "properties": {
            "title": {"type": ["string", "null"]},
            "unresolvedField": {"type": ["string", "null"]},
        },
        "required": ["title", "unresolvedField"],
        "additionalProperties": False,
    }

    with tempfile.TemporaryDirectory(prefix="internum-smoke-") as temp_dir:
        samples = _create_samples(Path(temp_dir))
        with TestClient(create_app()) as client:
            for sample in samples:
                response = _post_sample(client, api_key, sample, schema)
                if response.status_code != 200:
                    print(f"Smoke failed for {sample.name}: {response.status_code} {response.text}")
                    return 1
                body = response.json()
                data = body.get("data")
                meta = body.get("meta")
                if not isinstance(data, dict) or not isinstance(meta, dict):
                    print(f"Smoke failed for {sample.name}: response missing data/meta: {body}")
                    return 1
                if "unresolvedField" not in data or data["unresolvedField"] is not None:
                    print(f"Smoke failed for {sample.name}: unresolvedField was not null: {data}")
                    return 1
                if "documentType" not in meta:
                    print(f"Smoke failed for {sample.name}: meta missing documentType: {meta}")
                    return 1
                print(
                    f"Smoke passed for {sample.name}: "
                    f"{json.dumps({'data': data, 'meta': meta}, sort_keys=True)}"
                )

    return 0


def _post_sample(
    client: TestClient,
    api_key: str,
    sample: Path,
    schema: dict[str, Any],
):
    with sample.open("rb") as file:
        return client.post(
            "/v1/parse",
            headers={"X-API-Key": api_key},
            data={
                "schema": json.dumps(schema),
                "additionalContext": "The unresolvedField is intentionally absent.",
            },
            files={"file": (sample.name, file, "application/octet-stream")},
        )


def _create_samples(temp_dir: Path) -> Iterable[Path]:
    samples = [
        _create_pdf(temp_dir / "sample.pdf"),
        _create_scanned_pdf(temp_dir / "sample-scanned.pdf"),
        _create_docx(temp_dir / "sample.docx"),
        _create_html(temp_dir / "sample.html"),
        _create_xlsx(temp_dir / "sample.xlsx"),
        _create_png(temp_dir / "sample.png"),
    ]
    samples.extend(_create_legacy_office_samples(temp_dir))
    return samples


def _create_pdf(path: Path) -> Path:
    import fitz

    document = fitz.open()
    page = document.new_page()
    page.insert_text((72, 72), "Title: Smoke PDF")
    document.save(path)
    document.close()
    return path


def _create_docx(path: Path) -> Path:
    from docx import Document

    document = Document()
    document.add_paragraph("Title: Smoke DOCX")
    document.save(path)
    return path


def _create_xlsx(path: Path) -> Path:
    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    sheet["A1"] = "Title"
    sheet["B1"] = "Smoke XLSX"
    workbook.save(path)
    return path


def _create_html(path: Path) -> Path:
    path.write_text("<!doctype html><html><body><h1>Title: Smoke HTML</h1></body></html>")
    return path


def _create_png(path: Path) -> Path:
    from PIL import Image, ImageDraw

    image = Image.new("RGB", (500, 140), color="white")
    draw = ImageDraw.Draw(image)
    draw.text((20, 50), "Title: Smoke PNG", fill="black")
    buffer = io.BytesIO()
    image.save(buffer, format="PNG")
    path.write_bytes(buffer.getvalue())
    return path


def _create_scanned_pdf(path: Path) -> Path:
    import fitz

    image_path = path.with_suffix(".png")
    _create_png(image_path)
    document = fitz.open()
    page = document.new_page(width=500, height=140)
    page.insert_image(page.rect, filename=str(image_path))
    document.save(path)
    document.close()
    return path


def _create_legacy_office_samples(temp_dir: Path) -> list[Path]:
    binary = _libreoffice_binary()
    if binary is None:
        print("DOC/XLS smoke samples skipped: LibreOffice binary is unavailable.")
        return []

    samples: list[Path] = []
    xlsx_source = _create_xlsx(temp_dir / "legacy-source.xlsx")
    xls_path = _convert_office_sample(binary, xlsx_source, "xls", temp_dir)
    if xls_path is not None:
        samples.append(xls_path.rename(temp_dir / "sample.xls"))

    docx_source = _create_docx(temp_dir / "legacy-source.docx")
    doc_path = _convert_office_sample(binary, docx_source, "doc", temp_dir)
    if doc_path is not None:
        samples.append(doc_path.rename(temp_dir / "sample.doc"))
    else:
        print("DOC smoke sample skipped: LibreOffice could not create a DOC sample.")

    return samples


def _convert_office_sample(
    binary: str,
    source: Path,
    target_extension: str,
    outdir: Path,
) -> Path | None:
    completed = subprocess.run(
        [
            binary,
            "--headless",
            "--nologo",
            "--nolockcheck",
            "--nodefault",
            "--nofirststartwizard",
            "--convert-to",
            target_extension,
            "--outdir",
            str(outdir),
            str(source),
        ],
        capture_output=True,
        timeout=30,
        check=False,
    )
    output = source.with_suffix(f".{target_extension}")
    if completed.returncode != 0 or not output.exists():
        print(f"{target_extension.upper()} smoke sample skipped: LibreOffice conversion failed.")
        return None
    return output


def _libreoffice_binary() -> str | None:
    configured = os.getenv("CORE_LIBREOFFICE_BINARY", "soffice")
    path = Path(configured)
    if path.is_absolute() or "/" in configured:
        return str(path) if path.exists() else None
    return shutil.which(configured)


if __name__ == "__main__":
    raise SystemExit(main())
